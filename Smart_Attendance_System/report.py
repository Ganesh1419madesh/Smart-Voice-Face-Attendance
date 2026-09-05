"""
report.py - Generate Excel attendance reports using openpyxl
"""

import os
from datetime import date, datetime, timedelta
import openpyxl
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

from config import REPORTS_DIR, PRIMARY_COLOR, ACCENT_COLOR
from database import (get_attendance_range, get_all_students,
                      get_attendance_by_student)


# ── Colour helpers ────────────────────────────────────────────────────────────
_HDR_FILL    = PatternFill("solid", fgColor="1A237E")
_SUB_FILL    = PatternFill("solid", fgColor="283593")
_ALT_FILL    = PatternFill("solid", fgColor="E8EAF6")
_GREEN_FILL  = PatternFill("solid", fgColor="C8E6C9")
_RED_FILL    = PatternFill("solid", fgColor="FFCDD2")
_YELLOW_FILL = PatternFill("solid", fgColor="FFF9C4")

_THIN_BORDER = Border(
    left=Side(style="thin",   color="BDBDBD"),
    right=Side(style="thin",  color="BDBDBD"),
    top=Side(style="thin",    color="BDBDBD"),
    bottom=Side(style="thin", color="BDBDBD"),
)

_WHITE_FONT = Font(color="FFFFFF", bold=True, name="Calibri", size=11)
_BOLD_FONT  = Font(bold=True, name="Calibri", size=11)
_NORM_FONT  = Font(name="Calibri", size=10)


def _apply_header(ws, row: int, columns: list[str]):
    for col_idx, title in enumerate(columns, start=1):
        cell              = ws.cell(row=row, column=col_idx, value=title)
        cell.font         = _WHITE_FONT
        cell.fill         = _HDR_FILL
        cell.alignment    = Alignment(horizontal="center", vertical="center")
        cell.border       = _THIN_BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_len = max(
            (len(str(cell.value)) if cell.value else 0) for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(
            max_len + 4, 40
        )


def _status_fill(status: str):
    s = (status or "").lower()
    if "present" in s:
        return _GREEN_FILL
    if "late" in s:
        return _YELLOW_FILL
    return _RED_FILL


# ── Report generators ─────────────────────────────────────────────────────────

def generate_daily_report(target_date: str | None = None) -> str:
    """
    Generate a daily attendance report for a specific date.
    Returns the saved file path.
    """
    if target_date is None:
        target_date = date.today().isoformat()

    records  = get_attendance_range(target_date, target_date)
    students = get_all_students()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Daily Report"

    # Title row
    ws.merge_cells("A1:I1")
    title_cell           = ws["A1"]
    title_cell.value     = f"Daily Attendance Report — {target_date}"
    title_cell.font      = Font(bold=True, size=14, color="FFFFFF", name="Calibri")
    title_cell.fill      = _HDR_FILL
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Sub-header
    ws.merge_cells("A2:I2")
    ws["A2"].value     = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws["A2"].font      = Font(italic=True, size=9, color="757575", name="Calibri")
    ws["A2"].alignment = Alignment(horizontal="right")

    # Column headers
    headers = ["#", "Student ID", "Name", "Department", "Date",
               "Time In", "Time Out", "Status", "Photo"]
    _apply_header(ws, 3, headers)
    ws.row_dimensions[3].height = 20

    # Attendance map
    att_map = {r["student_id"]: r for r in records}

    row_num = 4
    for idx, student in enumerate(students, start=1):
        sid = student["student_id"]
        rec = att_map.get(sid)
        fill = _ALT_FILL if idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")

        values = [
            idx,
            sid,
            student["name"],
            student["department"],
            rec["date"] if rec else target_date,
            rec["time_in"]  if rec else "—",
            rec["time_out"] if rec else "—",
            rec["status"]   if rec else "Absent",
            rec.get("photo_path", "") if rec else "",
        ]
        for col_idx, val in enumerate(values, start=1):
            cell           = ws.cell(row=row_num, column=col_idx, value=val)
            cell.font      = _NORM_FONT
            cell.border    = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill      = fill

        # Colour status cell
        status_cell      = ws.cell(row=row_num, column=8)
        status_cell.fill = _status_fill(values[7])
        photo_path = rec.get("photo_path") if rec else None
        if photo_path:
            image_path = os.path.join(os.path.dirname(__file__), "photos", photo_path)
            if os.path.isfile(image_path):
                image = ExcelImage(image_path)
                image.width = 120
                image.height = 90
                ws.add_image(image, f"I{row_num}")
                ws.row_dimensions[row_num].height = 72
        row_num += 1

    # Summary row
    present = sum(1 for r in records if r.get("status") in ("Present", "Late"))
    absent  = len(students) - present

    ws.merge_cells(f"A{row_num}:E{row_num}")
    ws.cell(row=row_num, column=1,
            value=f"Total Students: {len(students)}  |  Present: {present}  |  Absent: {absent}")
    ws.cell(row=row_num, column=1).font      = _BOLD_FONT
    ws.cell(row=row_num, column=1).fill      = _SUB_FILL
    ws.cell(row=row_num, column=1).font      = Font(bold=True, color="FFFFFF",
                                                     name="Calibri", size=10)
    ws.cell(row=row_num, column=1).alignment = Alignment(horizontal="left",
                                                          vertical="center")

    _auto_width(ws)
    ws.freeze_panes = "A4"

    os.makedirs(REPORTS_DIR, exist_ok=True)
    filepath = os.path.join(REPORTS_DIR, f"daily_{target_date}.xlsx")
    wb.save(filepath)
    print(f"[REPORT] Daily report saved: {filepath}")
    return filepath


def generate_range_report(start_date: str, end_date: str,
                           department: str = "") -> str:
    """Generate an attendance report for a date range. Returns file path."""
    records  = get_attendance_range(start_date, end_date, department)
    students = get_all_students()
    if department:
        students = [s for s in students if s["department"] == department]

    wb = openpyxl.Workbook()

    # ── Sheet 1: Detailed records ─────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Detailed Records"

    ws1.merge_cells("A1:H1")
    ws1["A1"].value     = (f"Attendance Report  |  {start_date} to {end_date}"
                           + (f"  |  Dept: {department}" if department else ""))
    ws1["A1"].font      = Font(bold=True, size=13, color="FFFFFF", name="Calibri")
    ws1["A1"].fill      = _HDR_FILL
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    headers = ["#", "Date", "Student ID", "Name", "Department",
               "Time In", "Time Out", "Status"]
    _apply_header(ws1, 2, headers)

    for idx, rec in enumerate(records, start=1):
        row  = idx + 2
        fill = _ALT_FILL if idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        vals = [idx, rec["date"], rec["student_id"], rec["name"],
                rec["department"], rec.get("time_in", "—"),
                rec.get("time_out", "—"), rec.get("status", "Present")]
        for c, v in enumerate(vals, 1):
            cell           = ws1.cell(row=row, column=c, value=v)
            cell.font      = _NORM_FONT
            cell.border    = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill      = fill
        ws1.cell(row=row, column=8).fill = _status_fill(vals[7])

    _auto_width(ws1)
    ws1.freeze_panes = "A3"

    # ── Sheet 2: Per-student summary ──────────────────────────────────────────
    ws2 = wb.create_sheet("Student Summary")
    ws2.merge_cells("A1:F1")
    ws2["A1"].value     = "Student Attendance Summary"
    ws2["A1"].font      = Font(bold=True, size=13, color="FFFFFF", name="Calibri")
    ws2["A1"].fill      = _HDR_FILL
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 28

    s_headers = ["Student ID", "Name", "Department", "Days Present",
                 "Days Absent", "Attendance %"]
    _apply_header(ws2, 2, s_headers)

    all_dates = set()
    d = datetime.strptime(start_date, "%Y-%m-%d").date()
    end = datetime.strptime(end_date, "%Y-%m-%d").date()
    while d <= end:
        all_dates.add(d.isoformat())
        d += timedelta(days=1)
    total_days = len(all_dates)

    att_by_student: dict[str, set] = {}
    for r in records:
        att_by_student.setdefault(r["student_id"], set()).add(r["date"])

    for idx, student in enumerate(students, start=1):
        sid     = student["student_id"]
        present = len(att_by_student.get(sid, set()))
        absent  = total_days - present
        pct     = (present / total_days * 100) if total_days else 0
        row     = idx + 2
        fill    = _ALT_FILL if idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")

        for c, v in enumerate([sid, student["name"], student["department"],
                                present, absent, f"{pct:.1f}%"], 1):
            cell           = ws2.cell(row=row, column=c, value=v)
            cell.font      = _NORM_FONT
            cell.border    = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill      = fill

        # Colour attendance % cell
        pct_cell      = ws2.cell(row=row, column=6)
        pct_cell.fill = _GREEN_FILL if pct >= 75 else (_YELLOW_FILL if pct >= 50
                                                        else _RED_FILL)

    _auto_width(ws2)
    ws2.freeze_panes = "A3"

    os.makedirs(REPORTS_DIR, exist_ok=True)
    safe_dept = department.replace(" ", "_") if department else "all"
    filepath  = os.path.join(
        REPORTS_DIR, f"report_{start_date}_to_{end_date}_{safe_dept}.xlsx"
    )
    wb.save(filepath)
    print(f"[REPORT] Range report saved: {filepath}")
    return filepath


def generate_student_report(student_id: str) -> str:
    """Generate a per-student attendance report. Returns file path."""
    from database import get_student
    student = get_student(student_id)
    if not student:
        raise ValueError(f"Student {student_id} not found.")

    records = get_attendance_by_student(student_id)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Report"

    # Header
    ws.merge_cells("A1:F1")
    ws["A1"].value     = f"Attendance Report — {student['name']} ({student_id})"
    ws["A1"].font      = Font(bold=True, size=13, color="FFFFFF", name="Calibri")
    ws["A1"].fill      = _HDR_FILL
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:F2")
    ws["A2"].value     = (f"Dept: {student['department']}  |  "
                          f"Course: {student.get('course', '—')}  |  "
                          f"Total Records: {len(records)}")
    ws["A2"].font      = Font(italic=True, size=9, name="Calibri")
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["#", "Date", "Time In", "Time Out", "Method", "Status"]
    _apply_header(ws, 3, headers)

    for idx, rec in enumerate(records, start=1):
        row  = idx + 3
        fill = _ALT_FILL if idx % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        vals = [idx, rec["date"], rec.get("time_in", "—"),
                rec.get("time_out", "—"), rec.get("method", "—"),
                rec.get("status", "Present")]
        for c, v in enumerate(vals, 1):
            cell           = ws.cell(row=row, column=c, value=v)
            cell.font      = _NORM_FONT
            cell.border    = _THIN_BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.fill      = fill
        ws.cell(row=row, column=6).fill = _status_fill(vals[5])

    _auto_width(ws)
    ws.freeze_panes = "A4"

    os.makedirs(REPORTS_DIR, exist_ok=True)
    filepath = os.path.join(REPORTS_DIR, f"student_{student_id}.xlsx")
    wb.save(filepath)
    print(f"[REPORT] Student report saved: {filepath}")
    return filepath


if __name__ == "__main__":
    today = date.today().isoformat()
    path  = generate_daily_report(today)
    print(f"Report: {path}")
